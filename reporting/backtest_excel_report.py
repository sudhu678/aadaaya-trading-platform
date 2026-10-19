"""Standalone, bounded XLSX export of public educational simulations."""
from __future__ import annotations
from io import BytesIO
from math import isfinite
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MAX_ROWS_PER_SHEET = 250_000

def _safe(value):
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        if not isfinite(value):
            raise ValueError("Report contains a non-finite number")
        return value
    if isinstance(value, str):
        # Prevent spreadsheet formula injection from untrusted data.
        if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + value
        return value[:32000]
    raise TypeError("Unsupported report cell value")

def _sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for i, row in enumerate(rows, 1):
        if i > MAX_ROWS_PER_SHEET:
            raise ValueError("Report exceeds maximum rows")
        ws.append([_safe(v) for v in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(wrap_text=True)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = min(30, max(14, len(headers[col-1]) + 3))
    return ws

def build_xlsx_report(result: dict) -> bytes:
    if not isinstance(result, dict) or not isinstance(result.get("results"), dict):
        raise ValueError("Invalid backtest result")
    wb = Workbook()
    wb.remove(wb.active)
    results = result["results"]
    if len(results) > 10:
        raise ValueError("Too many ticker results")
    _sheet(wb, "Summary", ["Ticker", "Initial equity", "Final equity", "Return %", "Closed trades", "Open shares"],
           ([ticker, record["initial_equity"], record["final_equity"], record["return_pct"], len(record["trades"]), record["open_shares"]]
            for ticker, record in results.items()))
    _sheet(wb, "Trades", ["Ticker", "Entry date", "Exit date", "Entry price", "Exit price", "Shares", "PnL", "Exit reason"],
           ([ticker, trade.get("entry_date"), trade.get("exit_date"), trade.get("entry"), trade.get("exit"), trade.get("shares"), trade.get("pnl"), trade.get("reason")]
            for ticker, record in results.items() for trade in record["trades"]))
    _sheet(wb, "Equity", ["Ticker", "Date", "Equity"],
           ([ticker, point.get("date"), point.get("equity")]
            for ticker, record in results.items() for point in record["equity_curve"]))
    _sheet(wb, "Assumptions", ["Item", "Description"], [
        ["Strategy", "Independent educational SMA crossover"],
        ["Capital", "Separate starting capital per ticker; NOT shared portfolio capital"],
        ["Execution", "Next-session open; conservative stop gap-through"],
        ["Exclusions", "No slippage, commissions, dividend/corporate-action reconciliation"],
        ["Orders", "Advisory-only; no real orders"],
        ["Disclaimer", result.get("disclaimer", "Educational research only")],
    ])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
