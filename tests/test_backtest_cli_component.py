import json
import pandas as pd
import pytest
from scripts.backtest_cli import run
from reporting.backtest_json_report import build_json_report

def sample_loader(ticker, start, end):
    dates = pd.date_range(start, end, freq='B')
    close = [100 + i * .5 for i in range(len(dates))]
    return pd.DataFrame({'Open': close, 'High': [v+1 for v in close],
        'Low': [v-1 for v in close], 'Close': close,
        'Volume': [10000] * len(close)}, index=dates)

def args(path, fmt='json'):
    return ['--tickers', 'AAPL', '--start', '2025-01-01', '--end', '2025-06-01',
            '--fast', '5', '--slow', '10', '--format', fmt, '--out', str(path)]

def test_cli_json(tmp_path):
    path = tmp_path / 'result.json'
    assert run(args(path), loader=sample_loader) == path
    report = json.loads(path.read_text())
    assert 'AAPL' in report['results']
    assert report['advisory_only'] is True

def test_cli_xlsx(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / 'result.xlsx'
    run(args(path, 'xlsx'), loader=sample_loader)
    assert load_workbook(path, read_only=True).sheetnames == ['Summary','Trades','Equity','Assumptions']

def test_cli_no_overwrite(tmp_path):
    path = tmp_path / 'result.json'; path.write_text('existing')
    with pytest.raises(FileExistsError): run(args(path), loader=sample_loader)
    assert path.read_text() == 'existing'

def test_cli_wrong_suffix(tmp_path):
    with pytest.raises(ValueError): run(args(tmp_path / 'result.txt'), loader=sample_loader)

def test_json_nonfinite_rejected():
    with pytest.raises(ValueError): build_json_report({'results': {'X': {'bad': float('nan')}}})

def test_cli_rejects_bad_symbol(tmp_path):
    values = args(tmp_path / 'report.json'); values[1] = '../SECRET'
    with pytest.raises(ValueError): run(values, loader=sample_loader)
