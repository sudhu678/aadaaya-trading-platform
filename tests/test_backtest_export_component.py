import io
import unittest
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from dashboard.api.main import app
from dashboard.api.backtest_routes import get_loader
from reporting.backtest_excel_report import build_xlsx_report
import pandas as pd

class ReportTests(unittest.TestCase):
    def setUp(self):
        def loader(ticker, start, end):
            idx = pd.date_range("2025-01-01", periods=100, freq="B")
            close = [100 + i / 10 for i in range(100)]
            return pd.DataFrame({"Open": close, "High": [v+1 for v in close], "Low": [v-1 for v in close], "Close": close}, index=idx)
        app.dependency_overrides[get_loader] = lambda: loader
        self.client = TestClient(app)
        self.body = {"tickers": ["ABC"], "start": "2025-01-01", "end": "2025-06-01"}
    def tearDown(self):
        app.dependency_overrides.clear()
    def test_xlsx_endpoint(self):
        response = self.client.post("/api/public-backtest/export", json=self.body)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("spreadsheetml", response.headers["content-type"])
        wb = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(wb.sheetnames, ["Summary", "Trades", "Equity", "Assumptions"])
        self.assertEqual(wb["Summary"]["A2"].value, "ABC")
    def test_rejects_invalid_request(self):
        response = self.client.post("/api/public-backtest/export", json={**self.body, "fast": 60, "slow": 20})
        self.assertEqual(response.status_code, 422)
    def test_formula_injection_escaped(self):
        result = {"results": {"=SUM(1,2)": {"initial_equity": 100, "final_equity": 100, "return_pct": 0, "trades": [], "open_shares": 0, "equity_curve": []}}}
        wb = load_workbook(io.BytesIO(build_xlsx_report(result)))
        self.assertEqual(wb["Summary"]["A2"].data_type, "s")
        self.assertTrue(wb["Summary"]["A2"].value.startswith("'"))
    def test_rejects_nonfinite(self):
        result = {"results": {"ABC": {"initial_equity": float("inf"), "final_equity": 100, "return_pct": 0, "trades": [], "open_shares": 0, "equity_curve": []}}}
        with self.assertRaises(ValueError):
            build_xlsx_report(result)
